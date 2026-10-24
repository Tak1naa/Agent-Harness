from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path('deliverables');out.mkdir(exist_ok=True)
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='WenQuanYi Zen Hei';s._element.rPr.rFonts.set(qn('w:eastAsia'),'WenQuanYi Zen Hei');s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.15
d.styles['Normal'].paragraph_format.space_after=Pt(5)
for style in d.styles:
 for border in style.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 2'].font.size=Pt(12)
header=sec.header.paragraphs[0];header.text='北京交通大学大学生创新训练项目  ·  中期研究记录';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('2026年10月22日  /  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for i,row in enumerate(t.rows):
  for cell in row.cells:
   tcpr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9E9E9' if i==0 else 'FFFFFF');tcpr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcpr.append(borders)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2)
    for run in para.runs:run.font.size=Pt(9);run.bold=(i==0)
 return t


# Content follows only after complete, verified frozen results are available.
import json
run=Path('results/private/frozen-eval-20261022')
a=json.loads((run/'audit.json').read_text())
v=json.loads((run/'evidence-verification.json').read_text())
assert a['summary']['records']==216 and v['status']=='passed'
g={r['group']:r for r in a['summary']['groups']}
assert all(g[k]['denominator']==72 for k in 'ABC')
f=json.loads(Path('results/frozen-20261022/analysis.json').read_text())
pg={r['group']:r for r in f['primary_summary']['groups']}
dg=f['primary_diagnostics']; n=pg['A']['denominator']
d.add_heading('基于程序分析的代码智能体\n安全性测试 Harness 构建方法',0)
p('中期报告工作稿　2026年10月22日　依据已完成的冻结评测更新')
h('一、研究目标与范围调整')
p('本项目研究：在相同模型、工具和调用上限下，静态分析提供的告警位置及函数上下文，能否帮助模型生成更有效的Python安全测试。中期聚焦路径穿越、命令注入和SQL注入，以真实违规证据及缺陷版/修复版配对结果评估测试效果。')
p('目前进度落后于开题安排。保留立项题目，先完成单文件、统一入口的局部分析闭环。CFG与完整数据流、符号执行、自动修复、跨语言、完整智能体项目和200份样本尚未完成；本次范围调整仍需与导师确认。本文不预设提升比例，也不使用“首次提出”等未经论证的表述。')
h('二、实际完成的工作')
table(['工作','完成证据与边界'],[
 ['原型','Bandit/AST分析、受限测试生成、rootless容器执行、可信probe、JSON/CSV、CLI及MCP。'],
 ['工程验证','30对人工构造样本参考测试均配对成立；23项单元/协议测试和6项容器测试通过。'],
 ['模型实验',f'两轮开发试验与216次冻结试验已完成。主要比较保留{n}个A/B/C匹配块。'],
 ['真实生成数据','12个不同正常需求的模型输出全部保留，语法有效；动态执行、功能验收及双人标签仍待完成。'],
 ['框架接入','DeepSeek Harness固定版本SDK和MCP启动验证通过；带模型的完整工具回合尚未验证。'],
 ['执行地点','服务器曾复现三个代表案例；此次正式批量实验使用本机rootless Podman。']],[1.1,5.7])
p('已有20秒录像是人工参考测试的真实记录回放，用于展示证据链路。AI辅助代码与材料初稿，成员负责理解、核验和最终验收；本文不代替两人独立复核或真人计时记录。')
d.add_page_break()
h('三、方法与对照实验')
p('流程为Python代码、Bandit与AST分析、风险上下文、模型生成测试、隔离执行、证据汇总。Bandit提供已有规则告警；AST整理敏感调用所在函数、签名和导入。模型产出受限pytest测试输入，并可使用Hypothesis策略。目标入口、临时资源和安全断言由统一可信probe管理。当前未实现任意项目夹具自动生成。')
table(['组别','模型可见信息'],[['A 直接生成','正常需求、完整单文件源码、统一测试规范'],['B 告警引导','A组信息，加Bandit告警'],['C 上下文引导','B组信息，加AST定位的风险函数和依赖']],[1.3,5.5])
p('30组人工缺陷/修复配对样本中，6组用于开发、24组冻结评测，每类8组。每个样本每组重复3次，共216个trial、72个匹配块。样本是模板变体，不能称为30个独立真实项目；重复次数不增加独立样本量。修复版、标签和参考触发输入不进入生成提示，模型只根据缺陷版执行反馈修改测试。')
p('模型标识deepseek-flash，非思考模式，temperature=0，最大输出1800 Token。每次最多生成一次并修订两次，可提前成功停止。组序按冻结调度变化。三组的调用上限相同，输入Token和实际成本并不相同。提示、代码、数据和镜像均记录哈希，正式运行中未调整。上游模型别名不能等同于可见且永久不变的权重版本。')
p('路径穿越须读到临时目录外的随机哨兵，命令注入须创建无害标记，SQL注入须返回内存数据库中的PRIVATE行。同一最终测试在缺陷版违规且修复版通过，才计配对复现。普通异常、语法/导入错误、Mock及任意assert失败不能直接算漏洞。静态告警单独统计。')
p('执行使用无网络、非特权、只读根文件系统的rootless Podman容器，删除全部capability，限制CPU=1、内存256MiB、PID=64及超时。只挂载当前目标与测试，不传入密钥。隔离核验失败即停止。探针与目标同进程运行，尚不适合宣称能够抵抗恶意目标对判定器的篡改。')
p('DeepSeek Harness固定commit 687f282e8625c6b8d11ff58211a90524be2e4317，SDK 0.2.1rc2。正式实验使用直接API，以约束工具与重试预算；Harness是可替换编排层，不作为算法创新点。')
d.add_page_break()
h('四、冻结评测结果')
p(f'完成216/216次试验，原始响应、测试、执行目标/镜像哈希、配对结论及反馈链核验通过。出现{f["ambiguous_requests"]}次网络失败请求，无法确认服务端是否处理，全部保守计入调用次数并保留预算预留。超过3次上限的试验所在整个A/B/C块退出主要比较，共排除{len(f["excluded_matched_blocks"])}块，保留{n}块、{3*n}次试验。此处理规则在首次失败后、续跑前记录，并非原始预注册规则，因此同时披露全量结果。')
table(['指标','A 直接','B 告警','C 上下文'],[
 ['主要配对复现']+[f"{pg[k]['confirmed']}/{n} ({pg[k]['confirmed_rate']:.1%})" for k in 'ABC'],
 ['全量配对复现']+[f"{g[k]['confirmed']}/72" for k in 'ABC'],
 ['主要可运行']+[f"{pg[k]['runnable']}/{n}" for k in 'ABC'],
 ['首次缺陷违规']+[str(dg[k]['first_attempt_violation']) for k in 'ABC'],
 ['修复版误报/有效判定']+[f"{pg[k]['false_positives']}/{pg[k]['negative_valid_evaluations']}" for k in 'ABC'],
 ['修复版未判定']+[str(pg[k]['negative_unresolved']) for k in 'ABC'],
 ['成功保存API响应']+[str(dg[k]['responses']) for k in 'ABC'],
 ['输入/输出Token']+[f"{pg[k]['prompt_tokens']}/{pg[k]['completion_tokens']}" for k in 'ABC'],
 ['分支覆盖均值/分母']+[f"{pg[k]['mean_branch_coverage']:.1%} / {pg[k]['coverage_denominator']}" for k in 'ABC'],
 ['执行耗时合计(秒)']+[f"{pg[k]['duration_seconds']:.1f}" for k in 'ABC'],
 ['模型耗时合计(秒)']+[f"{pg[k]['model_duration_seconds']:.1f}" for k in 'ABC'],
 ],[1.9,1.6,1.6,1.6])
p(f'C相对A的主要复现率差为{100*(pg["C"]["confirmed_rate"]-pg["A"]["confirmed_rate"]):.1f}个百分点。这是当前模型、短程序模板与统一探针条件下的描述性结果，不能推断为普遍提升或统计显著。分支覆盖仅纳入存在分支且有有效覆盖记录的目标，公开分母；执行耗时与模型耗时分列，不含人工恢复等待。')
p(f'整个正式批次保存{a["api_responses"]}个响应，输入{a["usage"]["prompt_tokens"]}、输出{a["usage"]["completion_tokens"]} Token。按高峰未缓存输入2元/百万、输出8元/百万估算上界约{a["usage_cost_upper_cny"]:.4f}元，失败请求未知费用未包含在响应成本中，保守预留仍留在共享100元账本内。平台缓存、时段折扣及实际账单可能不同。')
d.add_page_break()
h('五、案例、失败与有效性限制')
p('成功案例：path-01人工参考测试输入../secret.txt，缺陷版直接拼接路径并读出公共目录外哨兵，修复版规范化路径并检查边界后拒绝访问。20秒配套录像展示该真实记录；人工参考成功率30/30只证明工程验证链路，不能作为模型实验效果。')
p('真实失败案例：path-10在repeat=0的A/B/C三组均未复现。目标只拒绝以../开头的输入，参考输入./../secret.txt仍可越界。生成测试尝试常见../、绝对路径及编码形式，却没有构造该有效绕过。记录为pass/pass不表示目标安全，而是该轮未找到违规。正式期间没有把参考输入补进提示，也未据此修改冻结样本。')
table(['漏洞类别','A配对复现','B配对复现','C配对复现'],[[{'path':'路径穿越','command':'命令注入','sql':'SQL注入'}[c]]+[f"{dg[k]['categories'][c]['confirmed']}/{dg[k]['categories'][c]['denominator']}" for k in 'ABC'] for c in ['path','command','sql']],[1.9,1.6,1.6,1.6])
p('主要限制包括：样本模板相似、单模型、统一函数入口、完整源码已提供给所有组、有限probe语法与同进程判定器。因此C组检验的是上下文显式组织的增益，尚不能证明大型项目依赖提取能力。安全性质通过也可能源于拒绝所有输入，必须另做正常功能验收。12份模型输出尚未动态验证和双人标注，不参与本轮人工配对结论。')
h('六、中期验收缺口与分工')
p('目前具备研究问题、可运行原型、三类可复现案例和一轮完整对照实验。提交前还需两人独立复核12份真实模型输出及正常功能，团队逐人重跑并解释案例，两轮计时演练，按学院模板填写身份字段并由导师确认范围调整。当前没有真人演练或双人标签完成证据。')
p('负责人负责研究问题、AST、集成和主讲；动态测试负责人负责probe、隔离、失败复现和录像；数据与报告负责人负责数据卡、标签、统计和材料。报告和PPT为可修改工作稿，学校模板、具体截止时间及署名信息仍待团队补齐。')
d.add_page_break()
h('七、中期后研究安排')
table(['阶段','工作与验收'],[
 ['第1—2周','补独立来源的真实缺陷与模型输出，建立来源/许可证/修复证据。功能与安全双重验收；按项目来源划分开发/评测。'],
 ['第3—5周','研究输入到敏感调用的局部数据依赖，提取边界校验与必要依赖。在较长、多函数代码中检验上下文是否有效。'],
 ['第6—8周','保留强直接生成基线，增加无LLM输入策略、去反馈/去依赖消融及第二模型；报告每次成功成本与失败类型。'],
 ['条件性扩展','仅当主要失败来自路径约束时，选择少量函数试验约束求解。自动修复在可信复现和功能回归之后，跨语言最后考虑。']],[1.1,5.7])
p('建议主线为“面向Python生成代码的风险上下文引导可验证安全测试”。下一步贡献应落到上下文提取与可验证判据的可解释增益，避免仅堆叠框架。置信区间应按独立样本/项目聚类；重复试验不当作独立样本，模板相似性需进一步检查。')
d.add_page_break()
h('八、主要参考文献')
refs=[
 '[1] Altmayer Pizzorno J, Berger E D. CoverUp: Effective High Coverage Test Generation for Python. FSE/PACMSE, 2025. DOI:10.1145/3729398.',
 '[2] Mathews N S, Nagappan M. Design choices made by LLM-based test generators prevent them from finding bugs. arXiv:2412.14137, 2024 (预印本).',
 '[3] Siddiq M L, Santos J C S. SecurityEval Dataset: Mining Vulnerability Examples to Evaluate Machine Learning-Based Code Generation Techniques. MSR4P&S, 2022. DOI:10.1145/3549035.3561184.',
 '[4] Barr E T, et al. The Oracle Problem in Software Testing: A Survey. IEEE TSE, 2015, 41(5):507–525. DOI:10.1109/TSE.2014.2372785.',
 '[5] Lemieux C, et al. CODAMOSA: Escaping Coverage Plateaus in Test Generation with Pre-trained Large Language Models. ICSE, 2023.',
 '[6] Kang S, Yoon J, Yoo S. Large Language Models are Few-shot Testers: Exploring LLM-based General Bug Reproduction. ICSE, 2023. arXiv:2209.11515.',
 '[7] Konstantinou M, Degiovanni R, Papadakis M. How well LLM-based test generation techniques perform with newer LLM versions? arXiv:2601.09695v2, 2026 (预印本).',
 '[8] Pearce H, et al. Asleep at the Keyboard? Assessing the Security of GitHub Copilot’s Code Contributions. IEEE S&P, 2022. arXiv:2108.09293.'
]
for ref in refs:p(ref)
p('完整链接与阅读顺序：docs/参考文献与阅读顺序.md；文献管理：docs/references.bib。Bandit与DeepSeek Harness作为工具资料另列。SecurityEval未直接用于本轮数据，CoverUp研究覆盖率而非本项目安全判据。开题所引约40%仅对应历史Copilot特定实验，不能当作现今模型的普遍漏洞率。')
d.save(out/'中期报告-20261022.docx')
print(out/'中期报告-20261022.docx')
