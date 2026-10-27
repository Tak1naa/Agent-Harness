from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path('deliverables');out.mkdir(exist_ok=True)
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='WenQuanYi Zen Hei';s._element.rPr.rFonts.set(qn('w:eastAsia'),'WenQuanYi Zen Hei');s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.15
d.styles['Normal'].paragraph_format.space_after=Pt(5)
for style in d.styles:
 for border in style.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 2'].font.size=Pt(12)
header=sec.header.paragraphs[0];header.text='北京交通大学大学生创新训练项目  ·  中期研究记录';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('2026年10月21日  /  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for i,row in enumerate(t.rows):
  for cell in row.cells:
   tcpr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9E9E9' if i==0 else 'FFFFFF');tcpr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcpr.append(borders)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2)
    for run in para.runs:run.font.size=Pt(9);run.bold=(i==0)
 return t

d.add_paragraph('基于程序分析的代码智能体\n安全性测试 Harness 构建方法',style='Title')
d.add_paragraph('大创项目中期报告',style='Subtitle')
p('本报告依据立项原件、当前代码及实际运行记录编写。中期已形成可运行原型并完成受控参考验证；真实模型对照实验尚未开展。前期进度落后于立项安排，现将范围收敛到Python单文件、统一函数入口和三类安全性质，优先建立可追溯的实验闭环。')
h('一 研究目标与范围调整')
p('研究问题是：在相同模型、工具和调用预算上限下，静态分析提供的风险位置及函数上下文，能否帮助模型生成更有效的Python安全测试？保留原立项题目，首阶段覆盖路径穿越、命令注入和SQL注入。符号执行、自动修复及跨语言支持列入后续，不预设提升幅度。')
p('项目以可重复的安全性质测试为目标。模型仅选择测试输入与受限Hypothesis策略，统一probe负责调用入口、临时资源和断言。任意项目夹具生成、跨文件污点分析和完整代码智能体轨迹评估目前均未完成。')
h('二 实际进展')
table(['交付项','当前状态','可追溯证据'],[
['分析与执行原型','已实现并验证','safety_harness/；containers/'],
['受控样本','30对；6开发、24评测','data/controlled/freeze.json'],
['参考测试验证','30/30配对成立','results/reference/*.json'],
['服务器复现','三类代表案例均成立','results/server/*.json'],
['SDK与MCP','无Key初始化与协议通过','results/dsh-integration.json'],
['真实模型实验','未开展，Key留空','暂无A/B/C效果结果'],
], [1.55,2.15,2.8])
p('AI辅助完成代码、资料整理和材料初稿；三位成员仍需理解方法、复核样本及判据并完成最终验收。报告不将AI辅助产物等同于已完成的双人标注或真人演练。')
d.add_page_break()
h('三 技术路线与实现边界')
p('流程依次为Python代码、Bandit与AST分析、风险上下文、模型测试生成、隔离执行、证据与统计。Bandit复用现有安全扫描规则；项目AST模块提取函数签名、导入、敏感调用及其所在函数，采用局部语法分析，不声称已实现完整污点传播。')
p('pytest管理测试执行，Hypothesis支持有界输入扩展。模型输出须通过AST语法白名单，只允许受限策略及probe.check调用。安全断言由可信实现统一定义，从而避免模型仅用assert False或Mock命中制造漏洞结论。正常功能正确性仍须单独核验。')
table(['安全性质','实际触发证据','修复版验收'],[
['路径穿越','返回公共目录外随机哨兵内容','同一测试不再越界'],
['命令注入','在容器临时区创建无害标记','标记不产生'],
['SQL注入','内存数据库返回PRIVATE行','私有行不返回'],
], [1.25,3.5,1.75])
p('主要判据为同一测试在缺陷版触发违规、修复版通过。普通异常、语法或导入失败、超时、虚假断言分别记录为失败；静态告警不与动态复现混算。生成阶段只反馈缺陷版运行信息，修复版与人工参考输入不进入模型提示。')
p('每次运行在rootless Podman容器中执行，使用无网络、只读根、非特权用户、全部capability删除、no-new-privileges和资源限制。限制为1 CPU、256MiB内存、64个进程，默认执行超时20秒。容器仅接收当前目标代码与测试，不获得API密钥或个人目录。启动前检查隔离参数，失败立即停止，无宿主机回退路径。')
p('服务器经SSH验证支持rootless Podman。镜像仓库连接失败后，通过现有SSH传输本地已验证OCI镜像，未修改网络或管理配置。三类缺陷／修复案例在服务器重复成立，本地与服务器镜像ID一致。')
p('DeepSeek Harness采用固定版本SDK和官方MCP stdio插件，暴露analyze_project、run_harness、summarize_run。SDK初始化与独立MCP工具调用已验证，带模型完整回合尚未验证。主实验保留直接API入口，以固定调用上限和预算。上游框架属于复用组件，项目贡献在上下文组织、探针、隔离证据与实验设计。')
d.add_page_break()
h('四 实验设计与当前结果')
table(['组别','模型可见信息','共同约束'],[
['A','需求、源码、统一测试规范','同模型、温度0、1800输出Token'],
['B','A 加 Bandit 告警','生成1次、反馈修订最多2次'],
['C','B 加 AST 风险函数与依赖','只改测试，每样本重复3次'],
], [0.6,3,2.9])
p('受控数据30组缺陷／修复配对，每类10组，其中6组用于开发、24组冻结评测。样本均为AI辅助人工构造的模板变体，明确记录来源与哈希，不能计作模型生成成果。正式受控评测共216个trial，最多648次生成调用。另用12个正常需求生成代码，每类4个，保留全部结果，由两名成员核验标签，未知项单列，两类数据分别报告。')
p('在每个样本及重复块内安排A/B/C，并随机化组序。三组看到同一份完整单文件源码，因此C组检验的是风险定位与上下文整理的增益。成功可以早停，报告实际调用和Token；相同的是预算上限。正式比较仅使用完整匹配块，失败项不能被静默删除。')
table(['参考测试集','可运行配对','确认复现','修复版触发'],[
['开发集','6/6','6/6','0/6'],['冻结评测集','24/24','24/24','0/24'],['合计','30/30','30/30','0/30']], [2,1.5,1.5,1.5])
p('上述结果来自人工参考测试，只证明当前受控样本和执行判据可用，不能推断模型效果。评测24例中仅5例含分支，最终参考测试平均分支覆盖为50%；其余19例无分支，不进入该均值。当前模型调用为0，暂无真实A/B/C胜负、提升率或显著性结论。')
p('路径案例path-01用../secret.txt触发缺陷版返回私有内容，修复版路径边界检查拒绝访问。普通断言失败的分类测试被记为harness_error，死循环容器测试按timeout终止；这些失败不会被记为漏洞。单元与协议测试21项、真实容器测试5项均通过，测试记录分别保存在JUnit文件中。')
p('主指标包括可运行率、缺陷复现率、安全版本误报率、被测代码分支覆盖、执行耗时及Token。每项公布分母、失败数与缺失值。JSON保留响应、执行与配对证据，CSV由日志重建。默认API规划总上限100元，先开发集小批量测成本；密钥仍为空，预算不是已发生费用。')
d.add_page_break()
h('五 当前问题与应对')
p('进度滞后的主要应对是缩小实现边界并优先完成可验证闭环。当前仅支持单文件和统一handle入口，手工模板短小且相似，开发／评测间存在模式相似风险。样本数量及来源限制了外推能力，后续引入独立项目并按来源或模板划分。无提升或负提升也须如实报告。')
p('安全性质通过不等于功能正确：拒绝全部输入可能通过安全测试。生成代码核验应补充正常功能用例。同进程目标代码若刻意篡改probe或运行报告，当前实现不能给出完全可信的对抗证明；后续可独立证据采集边界。容器隔离也不等于抵御未知内核漏洞。')
h('六 后续排期与分工')
table(['时段','工作与验收'],[
['D1—2','环境、样本规范与一个人工案例'],
['D3—5','闭环与框架接入，接入最多两个人日'],
['D6—8','三类案例、日志接口、冻结评测集'],
['D9—11','开发集A/B/C试跑、标签与成本核验'],
['D12—14','正式实验、复核、报告与录像更新'],
['D15—18','七页PPT、问答及两轮计时演练'],
], [1.05,5.45])
p('负责人负责研究问题、AST、框架集成与主讲；动态测试负责人负责测试生成、隔离、反馈和演示；数据与报告负责人负责样本、标签、相关工作、统计及材料。排期从团队确认启动日算起，收到学院模板和截止日期后调整。')
p('答辩共7页，页预算依次25、35、55、45、75、35、30秒，总计300秒。第4页配20秒静音真实记录回放，不依赖现场API。讲稿约900—1100汉字；真人两轮计时尚未开展，目标4分40秒至4分50秒，预留切页缓冲。')
h('资料与证据索引')
for text in ['[1] 项目开题报告，11页，原始PDF已阅读，含个人信息的原件不随公开仓库分发。',
'[2] Bandit官方文档：https://bandit.readthedocs.io/en/latest/ 。用于静态扫描方法与复用边界。',
'[3] pytest：https://docs.pytest.org/en/stable/ ；Hypothesis：https://hypothesis.readthedocs.io/en/latest/ 。',
'[4] DeepSeek Harness：https://github.com/deepseek-ai/deepseek-harness 。固定SDK 0.2.1rc2，源码提交见集成记录。',
'[5] SecurityEval：https://github.com/s2e-lab/SecurityEval 。仅作后续数据来源线索，未计入当前样本。',
'[6] 运行证据：results/reference/、results/server/、results/unit-tests.xml、results/container-tests.xml。资料核查日期2026-10-21。']:
 para=d.add_paragraph(text);para.paragraph_format.space_after=Pt(3)
 for run in para.runs:run.font.size=Pt(8)
d.save(out/'中期报告.docx')
print(out/'中期报告.docx')
