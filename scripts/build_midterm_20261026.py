"""从已有日志重建中期材料；不运行目标代码或调用模型。"""
from pathlib import Path
import json
import re
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables/中期答辩-20261026'
BUILD = ROOT / 'build/midterm-20261026'
OUT.mkdir(exist_ok=True, parents=True)
BUILD.mkdir(exist_ok=True, parents=True)
load = lambda p: json.loads((ROOT / p).read_text())
f = load('results/frozen-20261022/analysis.json')
b = load('results/benchmark-20261025/analysis.json')
c = load('results/private/client/validation-20261025.json')
pg = {r['group']: r for r in f['primary_summary']['groups']}
fg = {r['group']: r for r in f['full_summary']['groups']}
bg = {r['group']: r for r in b['within_budget_summary']['groups']}
assert f['evidence_verification']['status'] == b['evidence_verification']['status'] == 'passed'
assert [pg[g]['confirmed'] for g in 'ABC'] == [61,62,67]
assert [bg[g]['confirmed'] for g in 'ABC'] == [20,22,22]

slides = [
 dict(title='基于程序分析的代码智能体\n安全性测试 Harness 构建方法', time=25, short='研究问题',
      speech='各位老师好，我们的项目是基于程序分析的代码智能体安全性测试Harness构建方法。中期聚焦一个问题：在相同模型和调用上限下，把风险位置与函数上下文提供给模型，能否帮助它生成更有效的Python安全测试？'),
 dict(title='先把三类风险验证清楚', time=35, short='范围与判据',
      speech='目前进度落后于开题安排，我们先收敛到路径穿越、命令注入和SQL注入。静态工具提示哪里可疑，但是否真的违规，还要执行验证。比如读到了指定目录外的内容，才算路径穿越证据。先做局部分析和动态验证，符号执行、自动修复与跨语言延期。'),
 dict(title='程序如何完成一次安全测试', time=55, short='技术路线',
      speech='程序先读取代码，用Bandit扫描，再用AST整理风险函数、签名和导入。模型根据这些信息生成测试输入，程序把测试放进无网络、非特权容器执行。统一探针准备临时文件或数据库，并检查真实违规。没有成功时，只把缺陷版反馈给模型，最多修改两次。最后保存测试、执行日志和哈希，自动汇总统计。我们主要实现上下文组织、配对判定、受控执行和证据记录。正式实验直接调用API，官方Harness SDK目前只验证了初始化和MCP启动。'),
 dict(title='同一测试，核对缺陷版与修复版', time=45, short='配对复现',
      speech='这里用人工参考测试演示路径穿越。输入点点斜杠secret.txt后，缺陷版直接拼接路径，读到了目录外的随机哨兵。修复版先规范化路径并校验边界，同一测试不再触发。现在播放二十秒的真实运行记录回放。这验证的是执行和判定链路，模型生成测试的效果看下一页。'),
 dict(title='冻结批次 C 较高，补充试跑 B 与 C 持平', time=75, short='实验与失败',
      speech='冻结评测使用二十四个样本，每组重复三次，共二百一十六次试验。A直接给源码，B增加告警，C再增加上下文。网络失败也计入调用上限，排除三个超预算匹配块后，每组保留六十九次，分别配对复现六十一、六十二和六十七次。补充试跑每组二十四次，结果是二十、二十二和二十二，B与C持平。因此只能说当前批次观察到差异，不能说上下文总能提升。失败案例中，三组某轮都漏掉了点斜杠加点点斜杠的绕过方式。测试通过只表示这次没找到违规；重复试验也不等于更多独立程序。报告同时保留失败数量、修复版误报、覆盖率和成本，未知情况单列，不只展示成功结果。'),
 dict(title='原型已有客户端，研究验收仍需补齐', time=35, short='进展与分工',
      speech='现在已有会话式客户端，可以读代码、查看分析过程、运行已知案例，再展示修改差异供用户确认。真实模型调用和路径穿越配对已验证，客户端采用项目自有Agent。负责人负责方法、集成和主讲，另外两位分别负责动态测试、数据与材料。十二份真实生成代码已保留，但动态验收与双人真人标签仍待补齐，不能把AI初审当作真人复核。'),
 dict(title='下一步研究更真实的代码与更有效的上下文', time=30, short='后续研究',
      speech='接下来先补齐真实输出的正常功能、安全验收和双人复核。中期后引入独立来源样本，研究输入到敏感调用的局部依赖；通过去上下文、去反馈等对照，检查增益来自哪里，并比较成功成本。只有主要失败确实来自路径约束时，再试验有界符号执行。谢谢各位老师。'),
]
chars = len(re.findall(r'[\u4e00-\u9fff]', ''.join(s['speech'] for s in slides)))
assert 900 <= chars <= 1100, chars
assert sum(s['time'] for s in slides) == 300
content = dict(slides=slides, chinese_characters=chars,
               frozen=[pg[g] for g in 'ABC'], supplement=[bg[g] for g in 'ABC'])
(BUILD/'content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2))
shutil.copyfile(ROOT/'deliverables/演示回放.mp4',OUT/'演示回放.mp4')

script = ['# 五分钟逐页讲稿',f'正文约 {chars} 个汉字；7 页页预算共 300 秒，含第 4 页 20 秒录像。实际语速和切页须真人计时。','']
for i,s in enumerate(slides,1):
    script += [f'## 第 {i} 页　{s["short"]}　{s["time"]} 秒', '', s['speech'], '']
(OUT/'逐页讲稿.md').write_text('\n'.join(script))
plan = '''# PPT 大纲与 Presentation 方案

## 汇报主线

程序根据代码中的风险线索生成测试，在隔离环境中检查真实违规，并用缺陷版与修复版对照保留证据。现阶段已跑通系统并有小规模实验，研究重点是上下文的效果与适用边界。

## 七页大纲与节奏

| 页 | 主要内容 | 视觉设计 | 页预算 | 累计时间 |
|---|---|---|---:|---:|
| 1 | 原题目与一个研究问题 | 深色封面与测试证据示意 | 25 秒 | 0:25 |
| 2 | 三类风险、动态判据、延期范围 | 三行“风险—实际证据”对照 | 35 秒 | 1:00 |
| 3 | 程序的输入、处理、执行与输出 | 可编辑流程图，标明复用与自主实现 | 55 秒 | 1:55 |
| 4 | path-01 缺陷版与修复版 | 红绿分栏、相同输入、20 秒回放 | 45 秒 | 2:40 |
| 5 | A/B/C、两批结果、失败案例 | 原生柱状图和分母表，分批呈现 | 75 秒 | 3:55 |
| 6 | 客户端、分工、未完成项 | 真实客户端截图与验收状态 | 35 秒 | 4:30 |
| 7 | 中期后研究安排 | 分阶段研究路线与验收条件 | 30 秒 | 5:00 |

## 设计方案

16:9 横屏，深墨蓝、米白、青绿色构成主色；红色只用于违规，绿色用于配对通过，均同时带文字。封面和技术路线用深底，案例与结果用浅底。用流程、原生图表和真实截图承担说明，减少重复大字列表。保留原立项题目，不加入未经验证的创新性或提升承诺。图形与数据表可编辑，截图为真实界面取证。标题约 32–38 磅，主要说明约 19–23 磅，细节只放备注或报告。

## 如何讲

第 1 页直接说研究问题，不介绍大模型发展史。第 2 页用“读到目录外文件”等直白例子解释违规。第 3 页沿流程从左到右讲程序做了什么，明确 Bandit 是复用工具，贡献在上下文组织、执行判定和证据记录。第 4 页只讲一个完整例子；第 5 页先解释 A/B/C，再比较结果，最后说失败与局限。第 6 页展示客户端证明系统可以使用，不展开交互细节；第 7 页把后续工作落到可验收的实验。

转场提示：问题 → “怎样判断找到了漏洞？” → “程序具体怎么做？” → “看一次实际配对” → “模型是否因此更有效？” → “现在做到哪一步？” → “下一步验证什么？”

## 20 秒演示与现场备用

使用同目录 `演示回放.mp4`，它是既有人工参考测试真实记录的静音回放，不是新录制的客户端实操，也不是模型效果实验。视频作为独立文件交付，未嵌入 PPT。彩排时提前在播放器打开并暂停，放映第 4 页时切换播放，结束切回；切换方式和解码兼容性要在答辩电脑验证。

第 4 页前 12 秒说明输入与判据；中间 20 秒播放，可同步讲缺陷版/修复版；末尾 13 秒说明“这证明执行链路，模型效果看对照实验”。45 秒已包含录像，不额外加时。若视频无法播放，直接指向本页静态配对证据，用 15 秒说明结果后继续，现场不调用 API、不调试服务器。

## 彩排与交付检查

预算上限 300 秒；建议实际分配 23、32、52、43、70、33、30 秒，共 283 秒（4 分 43 秒），留下 17 秒用于切页和停顿。正文见《逐页讲稿.md》，与 PPT 备注一致。不要逐字念表格，分母、排除规则和额外指标留给追问。

至少两轮真人计时：第一轮完整讲并播放视频，记录每页结束时间；第二轮按删减后的表达重讲，目标 4:40–4:50。尚未执行真人演练，不将字符计数当作计时通过。提前核对投影可读性、字体替换、录像播放及切回位置。

## 重点问答

- **与 Bandit 有何区别？** Bandit 提供静态线索；本程序还生成并执行测试，用实际违规和修复版配对验证，静态告警不直接计漏洞。
- **自主贡献是什么？** 风险上下文组织、统一安全探针、隔离执行与证据记录、预算一致的对照流程，以及带修改审批的会话客户端；仍需验证方法的独立研究增益。
- **如何避免把测试失败当漏洞？** 检查目录外哨兵、无害执行标记和越权数据；语法错误、异常与 Mock 命中不能直接算漏洞。
- **为何暂缓符号执行？** 先完成可验证闭环，分析失败原因，再判断路径约束是否是主要障碍；当前没有完成该模块。
- **97.1%说明普遍有效吗？** 不能。它是冻结批次的 67/69，基于人工模板、单模型和重复试验；补充批次 B/C 均为 22/24。
- **这是官方 DeepSeek Harness 客户端吗？** 当前是本项目自有有界 Agent 客户端，直接调用 DeepSeek API；官方 SDK 只完成初始化与 MCP 启动验证。
- **12 份真实输出审完了吗？** 有外部初审反馈，但现有材料缺少完整动态验收及两位真人签署标签证据；未知仍保持未知。
- **公开 benchmark 跑通了吗？** 自建动态基准已跑；SecurityEval 仅做固定版本静态诊断，53/121 有任一线索，不是动态检出率或排行榜成绩。

## 证据边界

216 次冻结评测、72 次补充试跑、30 对人工参考验证、12 份真实生成输出以及 SecurityEval 静态诊断分别报告，不合并分母。原始数据来源及引用见中期报告末页。报告没有代填学校要求的身份、导师意见或签字；提交前需按学院正式模板核对。
'''
(OUT/'PPT大纲与Presentation方案.md').write_text(plan)

# 报告采用 A4 正式排版；显式分页使同类论证保持在一页。
d = Document(); sec=d.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.85);sec.bottom_margin=Cm(1.7)
sec.left_margin=sec.right_margin=Cm(2.1)
FONT='Noto Sans CJK SC'
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=d.styles[name];s.font.name=FONT;s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    s.font.color.rgb=RGBColor(0,0,0)
    for border in s.element.xpath('.//w:pBdr'):border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.line_spacing=Pt(16)
d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Title'].font.size=Pt(20)
d.styles['Heading 1'].font.size=Pt(15)
d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
d.styles['Heading 2'].font.size=Pt(11.5)
sec.header.paragraphs[0].text='北京交通大学大学生创新训练项目　中期研究报告'
sec.header.paragraphs[0].runs[0].font.size=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=2
foot.add_run('2026年10月26日　｜　').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
def p(t):return d.add_paragraph(t)
new_page=False
def h(t):
    global new_page
    para=d.add_heading(t,1)
    if new_page:para.paragraph_format.page_break_before=True
    new_page=False
    return para
def h2(t):return d.add_heading(t,2)
def page():
    global new_page
    new_page=True
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for cell,v in zip(t.rows[0].cells,headers):cell.text=v
    for row in rows:
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'19384A' if i==0 else ('F2F5F7' if i%2 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=Pt(13)
                for r in para.runs:
                    r.font.size=Pt(9);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    gap=d.add_paragraph();gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.line_spacing=Pt(4);gap.add_run().font.size=Pt(1)
    return t

d.add_paragraph('基于程序分析的代码智能体\n安全性测试 Harness 构建方法','Title')
p('中期报告工作稿　｜　研究与工程进展更新至 2026 年 10 月 26 日')
h('一 研究目标与实际进展')
p('我们已完成局部静态分析、模型生成测试、隔离执行与配对判定的原型，并完成一轮冻结对照实验和一轮补充试跑。现阶段结果表明风险上下文可能有帮助，但增益并不稳定；尚不能据此宣称适用于真实大型项目。新增会话式客户端提高了使用便利性，不作为方法有效性的额外实验样本。')
h2('研究问题和范围调整')
p('在相同模型、工具与调用上限下，Bandit告警及AST整理的风险函数上下文，能否帮助模型生成更有效的Python安全测试？中期覆盖路径穿越、命令注入、SQL注入，要求真实违规证据和缺陷版／修复版配对。')
p('项目进度落后于开题安排。保留原立项题目，将本阶段收敛到单文件、统一入口的局部分析；完整CFG与数据流、符号执行、自动修复、跨语言，以及开题中更大规模的生成代码验证尚未完成。开题预期的“提升30%”等目标不能写成实测结论。')
table(['工作','完成情况及证据边界'],[
 ['测试原型','Bandit与AST、受限pytest/Hypothesis测试、容器、probe、JSON/CSV、CLI与MCP接口已实现。'],
 ['研究实验','30对人工参考配对；216次冻结试验；72次补充试跑，三者分开统计。'],
 ['会话客户端','真实模型读文件、静态分析、path-01配对与修改提案已验证；应用修改需用户确认。'],
 ['真实模型输出','12份全部保留。已有外部初审反馈；完整动态验收与两名真人独立标签仍待补证。'],
 ['框架接入','主实验和客户端采用直接API；官方DeepSeek Harness SDK仅初始化及MCP启动验证。'],
 ['工程验证','2026-10-25记录：56项非容器Python、6项容器、5项JS测试通过；不等于67项端到端验收。']],[3.0,13.8])
p('AI用于辅助实现、审查与材料初稿；团队成员负责理解程序、独立复核、现场演示和最终验收。报告不代替真人签字、学院模板或导师意见。')

page();h('二 系统方法与判定依据')
p('流程为Python代码 → Bandit与AST → 风险上下文 → 模型生成测试 → 隔离执行 → 证据汇总。Bandit复用已有规则；AST按敏感调用名称定位所在函数，提取签名与导入。当前依赖信息主要是局部语法上下文，不等同于完整的数据流或污点分析。')
h2('测试如何生成和执行')
p('模型产生受限pytest测试及输入策略，可借助Hypothesis扩展边界；统一probe负责入口、临时资源与安全性质断言。每个试验最多一次生成、两次根据缺陷版反馈修订，只修改测试。模型不接收修复版、人工标签或参考触发输入。所有组均获得完整单文件源码，因此C组检验的是上下文显式组织的增益。')
table(['组别','模型获得的信息'],[['A 直接生成','需求、完整源码和统一测试规范'],['B 告警引导','A组信息 + Bandit告警'],['C 上下文引导','B组信息 + AST定位的风险函数、签名与导入']],[3.0,13.8])
table(['类型','动态证据','排除的错误判据'],[
 ['路径穿越','读到临时目录外的随机哨兵','仅产生路径字符串或异常不算'],
 ['命令注入','隔离环境出现无害执行标记','Mock命中、回显标记路径不算'],
 ['SQL注入','内存数据库返回越权PRIVATE行','查询报错或普通异常不算']],[3.0,7.0,6.8])
p('主要成功标准：同一最终测试在缺陷版触发违规、在修复版通过。语法错误、导入失败、超时与普通异常分别保留；测试通过只能说明本次未发现违规。正常功能另行验收，避免“拒绝所有输入”的实现被误写为合格。判定器设计参考测试预言问题[4]，生成反馈流程参考CoverUp[1]，漏洞发现能力不能用覆盖率替代[2]。')
h2('隔离与客户端边界')
p('目标在rootless、无网络、非特权、只读根文件系统容器中执行，删除capability并限制CPU、内存、PID与超时，不传入API密钥。隔离验证失败即停止，不回退宿主机。probe与目标同进程，目前不能宣称可抵抗恶意目标篡改判定器。')
p('客户端是项目自有有界Agent，通过DeepSeek API调度读文件、分析、已知案例配对等工具，不提供宿主机Shell。修改先展示diff，确认后检查原文SHA-256再写入。它是浏览器中的桌面式界面，尚非原生桌面安装包。官方SDK固定0.2.1rc2／commit 687f282e8625c6b8d11ff58211a90524be2e4317，未将启动验证写成完整模型回合验收。')

page();h('三 冻结评测设置与结果')
p('数据为30对人工构造缺陷／修复样本，每类10对；6对开发、24对冻结评测。每个样本A/B/C各重复3次，共216次试验、72个匹配块。样本含相似模板，216次试验不是216个独立程序。模型为deepseek-flash，非思考模式，temperature=0，最多1800输出Token，每试验最多3次请求；相同调用上限不意味着相同实际Token或费用。')
p('7次网络失败无法确认服务端是否处理，全部保守计为请求。超上限试验所在整个A/B/C匹配块退出主要比较，共排除3块，保留69块、207次试验。此规则在首次失败后、续跑前记录，并非原始预注册，故同时披露全量结果。原始响应、测试、目标／镜像哈希、反馈链核验通过，不能替代独立的安全判定审查。')
table(['指标','A 直接','B 告警','C 上下文'],[
 ['主要配对复现']+[f"{pg[k]['confirmed']}/69（{pg[k]['confirmed_rate']:.1%}）" for k in 'ABC'],
 ['全量配对复现']+[f"{fg[k]['confirmed']}/72" for k in 'ABC'],
 ['主要可运行']+[f"{pg[k]['runnable']}/69" for k in 'ABC'],
 ['缺陷最终未触发']+[str(pg[k]['statuses'].get('pass',0)) for k in 'ABC'],
 ['缺陷测试错误']+[str(pg[k]['statuses'].get('harness_error',0)) for k in 'ABC'],
 ['修复版误报／有效数']+[f"{pg[k]['false_positives']}/{pg[k]['negative_valid_evaluations']}" for k in 'ABC'],
 ['修复版未判定']+[str(pg[k]['negative_unresolved']) for k in 'ABC'],
 ['分支覆盖均值／分母']+[f"{pg[k]['mean_branch_coverage']:.1%} / {pg[k]['coverage_denominator']}" for k in 'ABC'],
 ['输入／输出Token']+[f"{pg[k]['prompt_tokens']} / {pg[k]['completion_tokens']}" for k in 'ABC'],
 ['执行耗时（秒）']+[f"{pg[k]['duration_seconds']:.1f}" for k in 'ABC'],
 ['模型耗时（秒）']+[f"{pg[k]['model_duration_seconds']:.1f}" for k in 'ABC'],
 ],[5.1,3.9,3.9,3.9])
p('C相对A高8.7个百分点，属于当前人工样本、单模型和统一probe条件下的描述性差异，不声称统计显著或普遍提升。分支覆盖仅纳入有分支且记录有效的目标，每组分母13；不能把76.9%当作全数据集总体覆盖。耗时为记录合计，不包含人工恢复等待。')
p(f"全批次保存{f['api_responses']}份响应，输入{f['usage']['prompt_tokens']:,}、输出{f['usage']['completion_tokens']:,} Token。按该批次记录的未缓存高峰单价计算，成功响应成本保守估计{f['usage_cost_upper_cny']:.4f}元；失败请求未知费用未包括，仍保留在共享100元预算预留中。历史估算及预留均不等于平台账单。")

page();h('四 补充评测与案例分析')
h2('补充动态试跑单独报告')
p('2026年10月25日对同一套24个评测样本，每组各运行一次，共72次试验；不是新增24个独立项目。预算与主要配置保持一致，未改核心提示和probe。107份成功响应及5次不确定请求均保留；每个试验仍在3次请求上限内，无匹配块剔除。')
table(['指标','A 直接','B 告警','C 上下文'],[
 ['配对复现']+[f"{bg[k]['confirmed']}/24（{bg[k]['confirmed_rate']:.1%}）" for k in 'ABC'],
 ['可运行']+[f"{bg[k]['runnable']}/24" for k in 'ABC'],
 ['未触发／测试错误','3 / 1','2 / 0','2 / 0'],
 ['修复版误报／有效数','0/24','0/24','0/24'],
 ['修复版未判定','0','0','0'],
 ],[5.1,3.9,3.9,3.9])
p('本批次B/C均为22/24，未观察到AST上下文相对告警的额外增益。分支覆盖仅5个有效记录，均为80%。输入91,545、输出21,271 Token；成功响应按历史单价估计0.353258元，账本预留增加2.540022元，二者不可混用。两批结果不能合并后再选择有利口径。')
h2('一个成功案例和一个失败案例')
p('path-01人工参考测试输入../secret.txt：缺陷版直接拼接路径，读到公共目录外哨兵；修复版规范化路径、检查边界后拒绝访问，同一测试通过。已有20秒视频回放这条真实记录。30/30人工参考配对证明执行链路可用，不能计为模型生成测试成绩。')
p('冻结批次path-10的repeat=0中三组均未复现。目标只拦截以../开头的输入，./../secret.txt仍可绕过；模型未构造有效输入。补充批次path-04、path-10也均未复现。失败可能来自输入选择及环境／判据配合，不能归结为目标安全。')
h2('SecurityEval仅完成静态诊断')
p('对官方固定提交6f4fb70f782c6d47b02ea24341e8ef8c1eb04a6a下121份Insecure_code分析：Bandit告警48份、AST敏感调用18份，任一线索53/121（43.8%），告警CWE与编号精确一致23/121。0个目标执行、0次模型调用。此处是线索覆盖，不是准确率、动态漏洞检出率或排行榜成绩。')
p('上游不安全标签未独立核验，缺少安全负例，入口与依赖资源尚未适配统一动态接口。SecurityEval[3]为外部样本来源参考；不能把该静态诊断与本项目配对复现分母相加。')

page();h('五 问题调整与后续计划')
h2('中期尚未完成的验收')
p('12份真实模型输出已保留，外部审查反馈提供了静态线索和待裁决问题；当前可核对材料中没有完整动态日志及两名真人签署标签。仍需在规定隔离环境中执行正常功能与安全测试，记录镜像、目标／测试哈希，完成两人独立标注并解决分歧。未知样本不删除、不当作安全。')
p('团队还需裁决.txt后缀约束、SQL前缀查询大小写语义及可信root内符号链接的威胁模型。这些是功能契约与攻击前提问题，不能凭静态可能性写成已确认漏洞。成员必须分别重跑并能解释代表案例，完成两轮真人计时。')
h2('与开题路径的关系及局限')
p('当前仍沿着“程序分析指导动态测试”的主线推进，已验证工程闭环与小规模对照；与原计划相比，覆盖面、分析深度和真实项目验证明显收窄。原计划的完整生成项目、200份规模、符号执行和自动修复并未实现。下一阶段优先解决真实样本与判据，再增加分析复杂度。')
p('有效性限制包括模板相似、单模型及可变模型别名、统一入口、完整源码已提供各组、局部名称匹配与同进程probe。当前结果既不证明能抵御任意恶意程序，也不证明C组在大型项目中稳定优于强基线。之后按独立来源／项目划分评测，避免把重复运行当作独立样本。')
table(['阶段','研究任务','验收依据'],[
 ['中期提交前','补12份动态和功能验收；两人独立标签；计时演练','可追溯日志、签署表与两轮计时'],
 ['中期后1—2周','扩充独立来源真实缺陷与模型输出，记录来源和许可','功能／安全双判据，按来源划分数据'],
 ['第3—5周','提取输入到敏感调用的局部数据依赖和边界校验','多函数代码中的上下文消融'],
 ['第6—8周','强直接生成基线、无LLM策略、去反馈消融及第二模型','成功率、误报、失败类别、每次成功成本'],
 ['条件性扩展','仅在路径约束成为主要失败原因时试验有界符号执行','小规模有界验证后再决定扩大']],[3.1,7.7,6.0])
p('分工：负责人负责问题、AST、集成与主讲；动态测试负责人负责probe、隔离和案例复现；数据与报告负责人负责数据、双人审查组织、统计与材料。自动修复应放在可靠复现和功能回归之后，跨语言最后考虑。')

page();h('六 证据索引与参考文献')
p('以下路径均相对项目仓库。公开统计可重建口径；原始私有日志保留证据，不随材料公开发布。新材料不包含立项原件、个人身份或API密钥。')
table(['证据','对应范围'],[
 ['results/frozen-20261022/analysis.json','216次全量／207次主要比较、排除规则、用量、核验范围'],
 ['results/benchmark-20261025/analysis.json','72次补充动态试跑、网络失败与预算'],
 ['results/benchmark-20261025/securityeval-static.json','121份外部静态诊断与来源'],
 ['results/private/client/validation-20261025.json','真实模型工具回合、path-01配对及客户端验证摘要'],
 ['safety_harness/；docs/客户端使用.md','分析、执行、probe、Agent与使用边界'],
 ['deliverables/演示回放.mp4','20秒人工参考测试真实记录回放']],[10.3,6.5])
p('研究定位：CoverUp[1]支持“覆盖反馈辅助生成”的方法背景；[2]提醒测试可运行和覆盖率不等于缺陷发现；SecurityEval[3]用于外部数据适配；测试预言综述[4]支持独立安全判据；CODAMOSA[5]与LIBRO[6]分别用于后续生成策略和缺陷复现对照设计。当前没有在相同数据和预算下复现这些系统，不能排名比较。')
refs=[
 '[1] Pizzorno J A, Berger E D. CoverUp: Effective High Coverage Test Generation for Python. PACMSE, 2025. DOI: 10.1145/3729398. https://arxiv.org/abs/2403.16218',
 '[2] Mathews N S, Nagappan M. Design choices made by LLM-based test generators prevent them from finding bugs. arXiv:2412.14137, 2024（预印本）. https://arxiv.org/abs/2412.14137',
 '[3] Siddiq M L, Santos J C S. SecurityEval Dataset: Mining Vulnerability Examples to Evaluate Machine Learning-Based Code Generation Techniques. MSR4P&S, 2022. DOI: 10.1145/3549035.3561184.',
 '[4] Barr E T, Harman M, McMinn P, et al. The Oracle Problem in Software Testing: A Survey. IEEE TSE, 2015, 41(5):507–525. DOI: 10.1109/TSE.2014.2372785.',
 '[5] Lemieux C, Inala J P, Lahiri S, Sen S. CODAMOSA: Escaping Coverage Plateaus in Test Generation with Pre-trained Large Language Models. ICSE, 2023. https://www.microsoft.com/en-us/research/publication/codamosa-escaping-coverage-plateaus-in-test-generation-with-pre-trained-large-language-models/',
 '[6] Kang S, Yoon J, Yoo S. Large Language Models are Few-shot Testers: Exploring LLM-based General Bug Reproduction. ICSE, 2023. https://arxiv.org/abs/2209.11515',
 '[7] PyCQA. Bandit documentation. https://bandit.readthedocs.io/（项目记录版本1.8.6）.',
 '[8] DeepSeek-AI. DeepSeek Harness. https://github.com/deepseek-ai/deepseek-harness（项目固定提交见正文）.',
]
for ref in refs:
    para=p(ref)
    for r in para.runs:r.font.size=Pt(9)
d.save(OUT/'中期报告.docx')
print(json.dumps({'output':str(OUT),'speech_chinese_characters':chars,'slides':7,'seconds':300},ensure_ascii=False))
